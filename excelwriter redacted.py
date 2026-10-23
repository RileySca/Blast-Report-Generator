# -*- coding: utf-8 -*-
"""
Created on Wed May 20 12:21:41 2026

"""

# excelwriter.py
from openpyxl import load_workbook
import shutil



def write_to_anchor(ws, cell_ref, value):
    """
    Writes to the top-left anchor of a merged range if the cell is merged.
    Otherwise writes directly.
    """
    for merged in ws.merged_cells.ranges:
        if cell_ref in merged:
            ws.cell(row=merged.min_row, column=merged.min_col).value = value
            return
    ws[cell_ref].value = value

def write_to_template(df, template_path, output_path, blast_date, Technician, Total, MaxDelay,new_blast_number):
    shutil.copy(template_path, output_path)
    
    wb = load_workbook(output_path)
    ws = wb.active

    # Header
    
    write_to_anchor(ws, "M4", blast_date)
    write_to_anchor(ws, "T4", df.iloc[0]["Time"])
    write_to_anchor(ws, "P15", Technician)
    write_to_anchor(ws, "F75", Total)
    write_to_anchor(ws, "Q75", Total)
    write_to_anchor(ws, "D78", MaxDelay)
    write_to_anchor(ws, "S78", MaxDelay)
    write_to_anchor(ws, "E4", new_blast_number)

    
    start_row = 25

    for i, row in df.iterrows():
        r = start_row + i

        ws[f"A{r}"] = i + 1  # Machine #

        write_to_anchor(ws, f"B{r}", row["Long_PPv"])
        write_to_anchor(ws, f"F{r}", row["Tran_PPv"])
        write_to_anchor(ws, f"J{r}", row["Vert_PPv"])
        write_to_anchor(ws, f"D{r}", row["Long_Freq"])
        write_to_anchor(ws, f"H{r}", row["Tran_Freq"])
        write_to_anchor(ws, f"L{r}", row["Vert_Freq"])

        #  Air dB with N/A fallback — MUST be indented here
        air_value = row["AirDB"]
        if air_value is None or air_value == "" or (
                isinstance(air_value, float) and air_value != air_value
                ):
            air_value = "N/A"

        write_to_anchor(ws, f"O{r}", air_value)
        write_to_anchor(ws, f"R{r}", row["Serial"])

        
    start_row = 18

        
    for i, row in df.iterrows():
        r = start_row + i

        ws[f"A{r}"] = i + 1  # Machine #

        write_to_anchor(ws, f"B{r}", row["Location"])
        write_to_anchor(ws, f"J{r}", row["Easting"])
        write_to_anchor(ws, f"O{r}", row["Northing"])
        write_to_anchor(ws, f"T{r}", row["Distance"])


    wb.save(output_path)